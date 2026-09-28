"""Unit tests for the pure logic: chunking, keyword tokens, faithfulness,
prompt building and experiment scoring. No Ollama needed.

    python -m unittest discover -s tests
"""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("LOCALBOT_DATA", tempfile.mkdtemp(prefix="localbot-test-"))

from localbot import documents, experiments, generation  # noqa: E402
from localbot.index import keyword_tokens  # noqa: E402


class ParagraphTests(unittest.TestCase):
    def test_numbered_headings_follow_the_document_numbering(self):
        lines = ["1. Setup", "Plug it in.", "2. Usage", "Run it.", "1. A list item inside usage", "3. Limits", "3.1 Power", "Max 5 W."]
        paras = documents._check_heading_numbers(documents._lines_to_paragraphs(lines))
        headings = [p["text"] for p in paras if p["level"]]
        self.assertEqual(headings, ["1. Setup", "2. Usage", "3. Limits", "3.1 Power"])

    def test_table_rows_stay_on_their_own_lines(self):
        paras = documents._lines_to_paragraphs(["Intro text.", "CMD    LEN    CHK", "0x02   2      0x03"])
        self.assertIn("CMD    LEN    CHK\n0x02   2      0x03", [p["text"] for p in paras])

    def test_markdown_headings_and_code_blocks(self):
        paras = documents._markdown_paragraphs("# Title\n\nText here.\n\n```\ncode\nmore\n```\n")
        self.assertEqual(paras[0]["level"], 1)
        self.assertTrue(any(p["text"].startswith("```") and "more" in p["text"] for p in paras))


class ChunkingTests(unittest.TestCase):
    def paras(self):
        out = []
        for s in range(1, 4):
            out.append({"text": f"{s}. Section {s}", "page": s, "level": 1})
            out += [{"text": f"Paragraph {s}.{i} " + "word " * 40, "page": s, "level": 0} for i in range(4)]
        return out

    def test_chunks_respect_size_and_record_sections(self):
        chunks = documents.chunk_paragraphs(self.paras(), 500, 80)
        self.assertTrue(all(len(c["text"]) <= 500 + 80 for c in chunks))
        self.assertTrue(all(c["section"] for c in chunks))
        self.assertEqual({c["page_start"] for c in chunks}, {1, 2, 3})

    def test_a_new_section_starts_a_new_chunk(self):
        chunks = documents.chunk_paragraphs(self.paras(), 2000, 0)
        self.assertEqual(len(chunks), 3)

    def test_overlap_must_be_smaller(self):
        with self.assertRaises(ValueError):
            documents.chunk_paragraphs(self.paras(), 100, 100)


class KeywordTests(unittest.TestCase):
    def test_identifiers_are_kept_whole_and_split(self):
        toks = keyword_tokens("Call Foo::BAR_BAZ with baudRate 0x1A3")
        for t in ("foo::bar_baz", "bar", "baz", "baudrate", "baud", "rate", "0x1a3"):
            self.assertIn(t, toks)
        self.assertNotIn("with", toks)


class FaithfulnessTests(unittest.TestCase):
    hits = [{"source": "My_Doc.pdf", "section": "", "text": "Ratio 12 /34 for ABCD- 5 at 0x1A3; baud 115200."}]

    def test_formatting_differences_are_not_flagged(self):
        c = generation.faithfulness("Per My_Doc, ratio 12/34, ABCD-5 at 0x1A3 [1].", self.hits, [], "q")
        self.assertEqual(c["unsupported"], [])
        self.assertTrue(c["cited"])

    def test_invented_values_and_citations_are_flagged(self):
        c = generation.faithfulness("It uses 0x1B7 and 9600 baud with reg_ctrl_2 [3].", self.hits, [], "q")
        self.assertEqual(c["unsupported"], ["0x1B7", "9600", "reg_ctrl_2"])
        self.assertEqual(c["bad_citations"], [3])


class PromptTests(unittest.TestCase):
    def test_passages_are_numbered_with_their_location(self):
        hits = [{"source": "a.pdf", "section": "2 Setup", "page_start": 3, "page_end": 3, "text": "Alpha."},
                {"source": "b.md", "section": "", "page_start": 0, "page_end": 0, "text": "Beta."}]
        p = generation.build_prompt("What?", hits)
        self.assertIn("[1] a.pdf, p. 3 — 2 Setup\nAlpha.", p)
        self.assertIn("[2] b.md\nBeta.", p)
        self.assertTrue(p.rstrip().endswith("say you don't know."))


class ScoringTests(unittest.TestCase):
    def test_short_tokens_need_word_boundaries(self):
        self.assertFalse(experiments.contains("please accept the change", "cc"))
        self.assertTrue(experiments.contains("use the cc command", "cc"))
        self.assertFalse(experiments.contains("sent 15 frames", "5"))

    def test_spacing_and_separators_are_ignored(self):
        self.assertTrue(experiments.contains("TX: AA 01 00 FF", "aa 01 00 ff"))
        self.assertEqual(experiments.fact_coverage("ttyS1 @ 9600 7E1", [["ttys1"], ["9600"], ["7e1"]]), 1)

    def test_labels(self):
        q = {"answerable": True, "facts": [["0x2a"]]}
        self.assertEqual(experiments.auto_label(q, "It is 0x2A [1].", 1.0), "correct")
        self.assertEqual(experiments.auto_label(q, "I don't know.", 0.0), "missed")
        self.assertEqual(experiments.auto_label({"answerable": False}, "The documents don't state it.", None), "correct")

    def test_question_set_validation(self):
        items, errors = experiments.parse_jsonl('{"id": "a", "question": "q?", "facts": [["x"]]}\n{"id": "b", "question": "q?"}\nnot json')
        self.assertEqual(len(items), 1)
        self.assertEqual(len(errors), 2)


class ClaimSplitTests(unittest.TestCase):
    def test_keeps_checkable_sentences_with_their_citations(self):
        answer = ("Answer:\nThe loader uses port A at 9600 baud [1]. It retries three times before giving up [2].\n"
                  "- The reset command is AA 01 00 FF [1].\n- Ok.\nI don't know what the default is.")
        claims = generation.split_claims(answer)
        self.assertEqual([c["cites"] for c in claims], [[1], [2], [1]])
        self.assertEqual(claims[2]["claim"], "The reset command is AA 01 00 FF.")
        self.assertTrue(all(c["text"] in answer for c in claims))  # needed for highlighting


class CarryOverTests(unittest.TestCase):
    def test_only_identical_answers_to_hand_graded_ones_inherit(self):
        from localbot import storage
        storage.init()
        graded = storage.log_interaction(origin="experiment", question="Q1?", answer="Same answer.")
        storage.query("UPDATE interactions SET grade = 'bad', tags = '[\"wrong\"]', reason = 'r', graded_at = ? WHERE id = ?",
                      (storage.now(), graded))
        run = storage.new_id()
        same = storage.log_interaction(origin="experiment", run_id=run, question="Q1?", answer="Same answer.")
        other = storage.log_interaction(origin="experiment", run_id=run, question="Q1?", answer="Different answer.")
        for qid, iid in (("a", same), ("b", other)):
            storage.query("INSERT INTO run_results VALUES (?, ?, ?, ?)", (run, qid, iid, "{}"))
        self.assertEqual(experiments.carry_over_grades(run), 1)
        row = storage.get_interaction(same)
        self.assertEqual((row["grade"], row["tags"], row["grade_source"]), ("bad", ["wrong"], f"carried:{graded}"))
        self.assertIsNone(storage.get_interaction(other)["grade"])
        # A carried grade is never used as a source for another carry.
        run2 = storage.new_id()
        again = storage.log_interaction(origin="experiment", run_id=run2, question="Q1?", answer="Same answer.")
        storage.query("INSERT INTO run_results VALUES (?, ?, ?, ?)", (run2, "a", again, "{}"))
        experiments.carry_over_grades(run2)
        self.assertEqual(storage.get_interaction(again)["grade_source"], f"carried:{graded}")


def _png(text="Reset command: AA 01 00 FF"):
    import io
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (400, 120), "white")
    ImageDraw.Draw(img).text((10, 50), text, fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


class FormatTests(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="localbot-formats-")

    def test_xlsx_rows_carry_their_column_names(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Ports"
        ws.append(["Port", "Baud", "Parity"])
        ws.append(["COM1", 9600, "none"])
        ws.append(["COM2", 115200.0, None])
        path = os.path.join(self.dir, "ports.xlsx")
        wb.save(path)
        texts = [p["text"] for p in documents.load_paragraphs(path)]
        self.assertEqual(texts, ["Sheet: Ports", "Port: COM1 · Baud: 9600 · Parity: none", "Port: COM2 · Baud: 115200"])

    def test_pptx_slides_tables_notes_and_pictures(self):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
        slide.shapes.title.text = "Setup"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        for r, row in enumerate([["Key", "Value"], ["Baud", "9600"]]):
            for c, v in enumerate(row):
                table.cell(r, c).text = v
        pic = os.path.join(self.dir, "p.png")
        open(pic, "wb").write(_png())
        slide.shapes.add_picture(pic, Inches(1), Inches(4))
        slide.notes_slide.notes_text_frame.text = "Mention the reset."
        path = os.path.join(self.dir, "deck.pptx")
        prs.save(path)
        paras = documents.load_paragraphs(path)
        self.assertEqual((paras[0]["text"], paras[0]["level"], paras[0]["page"]), ("Slide 1: Setup", 1, 1))
        self.assertIn("Key | Value\nBaud | 9600", [p["text"] for p in paras])
        self.assertTrue(any("image" in p for p in paras))
        self.assertEqual(paras[-1]["text"], "Speaker notes: Mention the reset.")

    def test_docx_pictures_are_found_in_place(self):
        import docx as docxlib
        d = docxlib.Document()
        d.add_paragraph("Before the picture.")
        pic = os.path.join(self.dir, "p.png")
        open(pic, "wb").write(_png())
        d.add_picture(pic)
        d.add_paragraph("After the picture.")
        path = os.path.join(self.dir, "doc.docx")
        d.save(path)
        kinds = ["image" if "image" in p else p["text"] for p in documents.load_paragraphs(path)]
        self.assertEqual(kinds, ["Before the picture.", "image", "After the picture."])

    def test_chunks_remember_their_pictures(self):
        paras = [{"text": "Intro text here.", "page": 1, "level": 0},
                 {"text": "[Picture, p. 1]\nReset command: AA 01 00 FF", "page": 1, "level": 0, "image_id": "0123456789abcdef"}]
        chunks = documents.chunk_paragraphs(paras, 800, 0)
        self.assertEqual(chunks[0]["images"], ["0123456789abcdef"])


class PictureCorrectionTests(unittest.TestCase):
    def test_correction_overrides_and_reverts(self):
        from localbot import vision
        image_id = vision.store(_png("Reset: AA 01"))
        with open(vision._cache_path(image_id), "w", encoding="utf-8") as f:
            f.write("model reading")
        self.assertEqual(vision.cached(image_id), "model reading")
        vision.set_correction(image_id, "  Reset: AA 01 00 FF  ")
        self.assertEqual((vision.cached(image_id), vision.is_corrected(image_id)), ("Reset: AA 01 00 FF", True))
        vision.clear_correction(image_id)
        self.assertEqual((vision.cached(image_id), vision.is_corrected(image_id)), ("model reading", False))


PAGE = """<html><head><title>Node</title><script>var x = 1;</script></head><body>
<nav class="sidebar"><a href="#">Classes</a></nav>
<div role="main">
<h1>Node<a class="headerlink" href="#node">¶</a></h1>
<p>Base class for all <em>scene</em> objects.</p>
<h2>Methods</h2>
<table><tr><th>Return</th><th>Method</th></tr><tr><td>void</td><td>queue_free ( )</td></tr></table>
<pre>func _ready():
    queue_free()</pre>
</div>
<footer>Built with a theme</footer></body></html>"""

SCRIPT = """import { x } from "./x.js";

export function add(a, b) {
  return a + b;
}

class Player {
  go() {}
}
"""


class ReaderTests(unittest.TestCase):
    def test_html_keeps_content_drops_site_chrome(self):
        paras = documents.paragraphs_from_bytes("node.html", PAGE.encode())
        text = "\n".join(p["text"] for p in paras)
        self.assertEqual([(p["level"], p["text"]) for p in paras[:2]], [(1, "Node"), (0, "Base class for all scene objects.")])
        self.assertIn("void | queue_free ( )", text)
        self.assertIn("```\nfunc _ready():\n    queue_free()\n```", text)
        for gone in ("Classes", "¶", "var x", "Built with"):
            self.assertNotIn(gone, text)

    def test_code_sections_follow_functions_and_classes(self):
        paras = documents.paragraphs_from_bytes("game.js", SCRIPT.encode())
        self.assertEqual([p["text"] for p in paras if p["level"]], ["add", "Player"])
        self.assertTrue(documents.looks_minified("var a=1;" * 500))
        self.assertFalse(documents.looks_minified(SCRIPT))


class ZipTests(unittest.TestCase):
    def test_docs_site_archive(self):
        import zipfile
        from localbot import index
        path = os.path.join(tempfile.mkdtemp(), "site.zip")
        with zipfile.ZipFile(path, "w") as z:
            z.writestr("site/classes/node.html", PAGE)
            z.writestr("site/classes/copy_of_node.html", PAGE)           # same bytes: duplicate
            z.writestr("site/game.js", SCRIPT)
            z.writestr("site/bundle.js", "var a=1;" * 500)               # minified
            z.writestr("site/_static/theme.js", SCRIPT + "// theme")     # site folder
            z.writestr("site/searchindex.js", "Search.setIndex({})")     # generated
            z.writestr("site/genindex.html", PAGE + " ")                 # generated
            z.writestr("__MACOSX/site/._node.html", "junk")
            z.writestr("site/logo.svg", "<svg/>")                        # not readable
            z.writestr("site/empty.md", "   ")
            z.writestr("site/notes/", "")
        chunks, info = index._chunk_file(path, 800, 100)
        self.assertEqual(info["files"], 2)
        self.assertEqual(info["skipped"], {"duplicate": 1, "minified or generated code": 1, "site or tool folder": 2,
                                           "generated site file": 2, "not a readable type": 1, "no readable text": 1})
        self.assertEqual(sorted({c["path"] for c in chunks}), ["site/classes/node.html", "site/game.js"])
        header = documents.chunk_header({"source": "site.zip", "path": "site/classes/node.html", "section": "Methods"})
        self.assertEqual(header, "site › site/classes/node.html — Methods")
        # A second read comes from the cache without opening the archive.
        again, _ = index._chunk_file(path, 800, 100, cached_only=True)
        self.assertEqual(len(again), len(chunks))


class KeywordIndexTests(unittest.TestCase):
    def test_postings_with_document_filter(self):
        from localbot import index

        class FakeCollection:
            name = "fake-keyword-test"

            def get(self, include):
                return {"ids": ["a", "b", "c"], "documents": ["queue_free frees the node", "the node tree", "unrelated text"],
                        "metadatas": [{"source": "godot.zip"}, {"source": "manual.pdf"}, {"source": "manual.pdf"}]}

        index._local.collection = FakeCollection()
        try:
            self.assertEqual([cid for _, cid in index.keyword_search("queue_free node", 5)], ["a", "b"])
            self.assertEqual([cid for _, cid in index.keyword_search("node", 5, sources={"manual.pdf"})], ["b"])
        finally:
            index._keyword_cache.pop("fake-keyword-test", None)
            index._local.collection = None


class JobTests(unittest.TestCase):
    def test_jobs_run_in_order_with_progress(self):
        import time
        from localbot import jobs

        def work(progress):
            progress(1, 2, "half")
            return {"ok": True}

        def fail(progress):
            raise ValueError("bad archive")

        a, b = jobs.submit("add", "A", work), jobs.submit("add", "B", fail)
        for _ in range(100):
            if jobs.get(b)["finished"]:
                break
            time.sleep(0.02)
        self.assertEqual((jobs.get(a)["status"], jobs.get(a)["result"], jobs.get(a)["step"]), ("done", {"ok": True}, "half"))
        self.assertEqual((jobs.get(b)["status"], jobs.get(b)["error"]), ("failed", "bad archive"))
        self.assertIsNone(jobs.active())


class ComparisonTests(unittest.TestCase):
    def test_comparison_words(self):
        for q in ("Compare Bellman-Ford with Dijkstra", "BFS vs DFS", "difference between A and B",
                  "how does A differ from B", "is A faster than B"):
            self.assertTrue(generation.COMPARE_RE.search(q), q)
        for q in ("Explain breadth first search", "What is the diff tool?", "list the vertices"):
            self.assertFalse(generation.COMPARE_RE.search(q), q)

    def test_follow_up_words(self):
        for q in ("How does it compare with Dijkstra?", "what about the second one", "why are they slow"):
            self.assertTrue(generation.REFERS_BACK_RE.search(q), q)
        for q in ("how does a binary tree differ from a hash map?", "explain iterators", "Is BFS complete?"):
            self.assertFalse(generation.REFERS_BACK_RE.search(q), q)

    def test_every_side_gets_passages(self):
        from localbot import retrieval
        from localbot.config import S

        def hit(i, score):
            return {"id": i, "score": score, "text": i}

        results = {  # the whole question only finds one side; below-cutoff ones aren't sent
            "compare A with B": [hit("a1", .9), hit("a2", .8), hit("ab", .4), hit("a3", .7)],
            "What is A?": [hit("a1", .95), hit("a4", .6)],
            "What is B?": [hit("b1", .9), hit("ab", .85), hit("b2", .3)],
        }
        original = retrieval.retrieve
        retrieval.retrieve = lambda q, **kw: (results[q], {"embed_ms": 1.0, "search_ms": 1.0, "rerank_ms": 1.0}, [0])
        try:
            with S.override({"top_k": 3, "rerank": True, "min_score_reranked": 0.5}):
                hits, timings, _ = retrieval.retrieve_comparison("compare A with B", ["A", "B"])
        finally:
            retrieval.retrieve = original
        self.assertEqual([h["id"] for h in hits[:4]], ["a1", "b1", "a2", "a4"])  # top_k 3 + 1 extra side
        self.assertEqual(hits[0]["found_for"], ["the question", "A"])
        self.assertEqual({h["id"] for h in hits[4:]}, {"ab", "a3", "b2"})  # logged, not sent
        with S.override({"rerank": True, "min_score_reranked": 0.5}):
            self.assertEqual([h["id"] for h in retrieval.sendable(hits)], ["a1", "b1", "a2", "a4"])
        self.assertEqual(timings["rerank_ms"], 3.0)

    def test_history_keeps_only_the_same_subject(self):
        from localbot import index
        from localbot.config import S
        common = {"algorithm", "graph", "shortest", "path"}
        original = index.distinctive_terms
        index.distinctive_terms = lambda text: {w for w in keyword_tokens(text) if len(w) > 2 and w not in common
                                                and w not in index.QUESTION_WORDS}
        turns = [{"question": "Explain BFS and compare with DFS", "metrics": {}},
                 {"question": "Explain the Bellman-Ford algorithm", "metrics": {}},
                 {"question": "what graph is used?", "metrics": {}}]
        try:
            with S.override({"focus_history": True}):
                kept = generation.relevant_turns(turns, "compare it with Dijkstra",
                                                 "Compare the Bellman-Ford algorithm with Dijkstra's algorithm")
                self.assertEqual(kept, [turns[1]])
                self.assertEqual(generation.relevant_turns(turns, "why?", "why"), [turns[2]])
            with S.override({"focus_history": False}):
                self.assertEqual(generation.relevant_turns(turns, "why?", "why"), turns)
        finally:
            index.distinctive_terms = original


class PortTests(unittest.TestCase):
    def test_busy_port_falls_back_with_a_reason(self):
        from localbot.server import open_listening_socket
        first, _ = open_listening_socket(0)
        port = first.getsockname()[1]
        second, notes = open_listening_socket(port)
        try:
            self.assertNotEqual(second.getsockname()[1], port)
            self.assertIn(str(port), notes[0])
        finally:
            first.close()
            second.close()


if __name__ == "__main__":
    unittest.main()
